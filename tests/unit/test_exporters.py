"""
Unit tests for packages/reports/exporters.py — pure functions, no database.
"""

from datetime import date
from decimal import Decimal

import pytest

from packages.reports.exporters import to_csv_bytes, to_excel_bytes, to_pdf_bytes

pytestmark = pytest.mark.unit

COLUMNS = ["name", "amount", "when"]
ROWS = [
    {"name": "Alpha", "amount": Decimal("123.456"), "when": date(2026, 1, 1)},
    {"name": "Beta", "amount": 99.5, "when": None},
]


def test_to_csv_bytes_includes_header_and_rows():
    result = to_csv_bytes(COLUMNS, ROWS)
    text = result.decode("utf-8-sig")
    lines = text.strip().splitlines()

    assert lines[0] == "name,amount,when"
    assert lines[1] == "Alpha,123.46,2026-01-01"
    assert lines[2] == "Beta,99.50,"


def test_to_csv_bytes_empty_rows_still_has_header():
    result = to_csv_bytes(COLUMNS, [])
    text = result.decode("utf-8-sig")
    assert text.strip() == "name,amount,when"


def test_to_excel_bytes_produces_valid_workbook():
    result = to_excel_bytes(COLUMNS, ROWS, sheet_name="My Report")
    assert result.startswith(b"PK")  # xlsx is a zip archive

    from openpyxl import load_workbook
    from io import BytesIO

    workbook = load_workbook(BytesIO(result))
    sheet = workbook["My Report"]
    assert [cell.value for cell in sheet[1]] == COLUMNS
    assert sheet.cell(row=2, column=1).value == "Alpha"


def test_to_excel_bytes_truncates_long_sheet_name():
    long_name = "x" * 50
    result = to_excel_bytes(COLUMNS, ROWS, sheet_name=long_name)

    from openpyxl import load_workbook
    from io import BytesIO

    workbook = load_workbook(BytesIO(result))
    assert len(workbook.sheetnames[0]) <= 31


def test_to_pdf_bytes_produces_valid_pdf():
    result = to_pdf_bytes("Test Report", COLUMNS, ROWS, subtitle="For the week")
    assert result.startswith(b"%PDF-")
    assert len(result) > 100
