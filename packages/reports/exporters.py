"""
Shared tabular export engine.

Every report in `modules.reports` (and any future module that wants to
let a user download a table) converts its rows through here rather than
each module reimplementing CSV/Excel/PDF writing. Input is always the
same shape — a list of column keys plus a list of row dicts — so the
export format is a pure function of that shape, independent of which
report produced it.
"""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def to_csv_bytes(columns: list[str], rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_stringify(row.get(col)) for col in columns])
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 CSV correctly


def to_excel_bytes(columns: list[str], rows: list[dict], sheet_name: str = "Report") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31] or "Report"  # Excel sheet name limit

    header_font = Font(bold=True)
    for col_index, column_name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=column_name)
        cell.font = header_font

    for row_index, row in enumerate(rows, start=2):
        for col_index, column_name in enumerate(columns, start=1):
            value = row.get(column_name)
            if isinstance(value, Decimal):
                value = float(value)
            sheet.cell(row=row_index, column=col_index, value=value if not isinstance(value, (dict, list)) else _stringify(value))

    for col_index, column_name in enumerate(columns, start=1):
        max_length = max(
            [len(column_name)] + [len(_stringify(row.get(column_name))) for row in rows]
        )
        sheet.column_dimensions[get_column_letter(col_index)].width = min(max_length + 2, 60)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf_bytes(title: str, columns: list[str], rows: list[dict], subtitle: str | None = None) -> bytes:
    buffer = io.BytesIO()
    page_size = landscape(letter) if len(columns) > 5 else letter
    doc = SimpleDocTemplate(buffer, pagesize=page_size, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()

    elements = [Paragraph(title, styles["Title"])]
    if subtitle:
        elements.append(Paragraph(subtitle, styles["Normal"]))
    elements.append(Paragraph("&nbsp;", styles["Normal"]))

    table_data = [columns] + [[_stringify(row.get(col)) for col in columns] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buffer.getvalue()
